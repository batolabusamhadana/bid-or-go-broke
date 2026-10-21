from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
wb=Workbook()
H=PatternFill('solid',fgColor='8B2BFF');IN=PatternFill('solid',fgColor='EFE5FF');F=Font(name='Arial',size=11);HF=Font(name='Arial',bold=True,color='FFFFFF')
th=Side(style='thin',color='B98CFF');BD=Border(left=th,right=th,top=th,bottom=th)
def sheet(ws,heads,rows,widths):
    ws.append(heads)
    for r in rows: ws.append(r)
    for c in ws[1]: c.fill=H;c.font=HF;c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=F;c.border=BD;c.alignment=Alignment(vertical='center',wrap_text=True)
    for i,w in enumerate(widths): ws.column_dimensions[chr(65+i)].width=w
    ws.freeze_panes='A2'
S=[('Amazon EC2','Compute',20),('Amazon S3','Object Storage',10),('AWS Lambda','Serverless Compute',15),('Amazon RDS','Relational Database',25),('Amazon DynamoDB','NoSQL Database',20),('Amazon CloudFront','Content Delivery',15),('Amazon Route 53','DNS',10),('Amazon Cognito','Login & Identity',10)]
ws=wb.active;ws.title='Start Here'
for r in [['BID OR GO BROKE · AWS SBG BZU · DEVFEST BZU'],[''],['HOW TO USE'],['1. Type team names in the Credits tab (purple-light cells, A3:A10).'],['2. After each auction round, fill the Purchase Log (winner, price, 2nd-copy team and price). Credits update automatically.'],['3. At the end, fill Scoreboard: must-haves covered (0-3), bonus count, overkill count, pitch (0-4). Totals and ranks are automatic.'],['4. Light purple cells = inputs. Everything else is a formula. Prices are fictional game values.'],['Scoring: 10 per must-have, +5 per bonus, -5 per overkill, +1 per 5 unspent CC (max 6), pitch 0-4. Incomplete builds (under 3 must-haves) cannot win.']]: ws.append(r)
ws['A1'].font=Font(name='Arial',bold=True,size=16,color='8B2BFF');ws['A3'].font=Font(name='Arial',bold=True,size=12)
for r in range(4,9): ws.cell(r,1).font=F
ws.column_dimensions['A'].width=130
ws=wb.create_sheet('Price List');sheet(ws,['Service','Category','Start (CC)','Min raise','Stock'],[[a,b,c,2,2] for a,b,c in S],[24,24,12,10,8])
ws=wb.create_sheet('Credits');sheet(ws,['Team','Start CC','Spent','Credits left','Cards owned'],[],[24,10,10,12,12])
ws.insert_rows(2);ws['A2']='Example: Team Orbit';ws['A2'].font=Font(name='Arial',italic=True,color='777777')
for r in range(3,11):
    ws.cell(r,1).fill=IN;ws.cell(r,2,100);ws.cell(r,3,f"=SUMIF('Purchase Log'!$C$3:$C$40,A{r},'Purchase Log'!$D$3:$D$40)+SUMIF('Purchase Log'!$E$3:$E$40,A{r},'Purchase Log'!$F$3:$F$40)")
    ws.cell(r,4,f'=IF(A{r}="","",B{r}-C{r})');ws.cell(r,5,f"=IF(A{r}=\"\",\"\",COUNTIF('Purchase Log'!$C$3:$C$40,A{r})+COUNTIF('Purchase Log'!$E$3:$E$40,A{r}))")
    for c in range(1,6): ws.cell(r,c).border=BD
    for c in range(2,6): ws.cell(r,c).font=F
ws['G1']='Rows: example on row 2 (ignored). Edit A3:A10.';ws['G1'].font=Font(name='Arial',size=9,italic=True)
ws=wb.create_sheet('Purchase Log');sheet(ws,['Round','Service','Winning team','Price paid','2nd-copy team','2nd price','Check'],[[i+1,S[i][0],None,None,None,None,None] for i in range(8)],[8,24,18,12,18,12,26])
ws.insert_rows(2);ws['A2']='Example';ws['B2']='Amazon S3';ws['C2']='Team Orbit';ws['D2']=16;ws['E2']='Team Nova';ws['F2']=14
for c in ws[2]: c.font=Font(name='Arial',italic=True,color='777777')
ws['B1'].value='Service'
for r in range(3,11):
    ws.cell(r,7,f'=IF(C{r}="","",IF(D{r}<VLOOKUP(B{r},\'Price List\'!$A$2:$C$9,3,FALSE),"Below start price!","OK"))')
    for c in (3,4,5,6): ws.cell(r,c).fill=IN
for r in range(11,31):
    for c in range(1,8): ws.cell(r,c).border=BD
    for c in (3,4,5,6): ws.cell(r,c).fill=IN
ws['A31']='Rows 11-30 are spare for extra games.';ws['A31'].font=Font(name='Arial',size=9,italic=True)
ws=wb.create_sheet('Scoreboard');sheet(ws,['Team','Scenario','Must-haves covered (0-3)','Bonus picks','Overkill items','Credits left','Pitch (0-4)','Must pts','Bonus pts','Overkill pts','Budget pts','TOTAL /50','Status','Rank'],[],[24,16,14,10,10,10,10,9,9,10,10,10,13,8])
ws.row_dimensions[1].height=45
for r in range(2,10):
    c=r+1
    ws.cell(r,1,f'=IF(Credits!A{c}="","",Credits!A{c})');ws.cell(r,6,f'=IF(A{r}="","",Credits!D{c})')
    for col in (2,3,4,5,7): ws.cell(r,col).fill=IN
    ws.cell(r,8,f'=10*N(C{r})');ws.cell(r,9,f'=5*MIN(2,N(D{r}))');ws.cell(r,10,f'=-5*MIN(2,N(E{r}))');ws.cell(r,11,f'=IF(A{r}="",0,MIN(6,INT(N(F{r})/5)))')
    ws.cell(r,12,f'=IF(A{r}="","",H{r}+I{r}+J{r}+K{r}+MIN(4,N(G{r})))')
    ws.cell(r,13,f'=IF(A{r}="","",IF(N(C{r})>=3,"Valid","INCOMPLETE"))')
    ws.cell(r,14,f'=IF(M{r}="Valid",COUNTIFS($M$2:$M$9,"Valid",$L$2:$L$9,">"&L{r})+1,"-")')
    for col in range(1,15): ws.cell(r,col).border=BD;ws.cell(r,col).font=F
dv=DataValidation(type='whole',operator='between',formula1='0',formula2='3');ws.add_data_validation(dv);dv.add('C2:C9')
dv2=DataValidation(type='whole',operator='between',formula1='0',formula2='4');ws.add_data_validation(dv2);dv2.add('G2:G9')
ws['A11']='Tie-break: more credits left, then best pitch (decide manually). Team names come from the Credits tab.';ws['A11'].font=Font(name='Arial',size=9,italic=True)
wb.save('../downloads/Bid_or_Go_Broke_Host_Sheet.xlsx')
